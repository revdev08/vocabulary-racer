"""Importer integration checks; run with Python + openpyxl."""
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

import openpyxl


class VocabularyImportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        (self.root / 'scripts').mkdir()
        (self.root / 'src/game/data').mkdir(parents=True)
        content = self.root / 'content/es-en'
        content.mkdir(parents=True)
        self.source = content / 'espanol-ingles.xlsx'
        self.script = self.root / 'scripts/import-vocabulary.py'
        shutil.copyfile(Path(__file__).resolve().parents[2] / 'scripts/import-vocabulary.py', self.script)
        book = openpyxl.Workbook()
        words = book.active
        words.title = 'Palabras'
        words.append(['Nivel', 'Tema', 'English', 'Español', 'Distractor 1', 'Distractor 2'])
        words.append([1, 'Acciones', 'look up', 'mirar up', 'sleep', 'walk'])
        words.append([1, 'Acciones', 'look up', 'mirar up', 'run', 'eat'])
        units = book.create_sheet('Niveles')
        units.append(['ID', 'Nombre', 'Objetivo'])
        units.append([1, 'Acciones', 'Aprender acciones'])
        book.save(self.source)

    def run_import(self, succeeds=True):
        result = subprocess.run([sys.executable, str(self.script)], capture_output=True, text=True)
        if succeeds:
            self.assertEqual(result.returncode, 0, result.stderr)
        else:
            self.assertNotEqual(result.returncode, 0)
        return result

    def catalog(self):
        text = (self.root / 'src/game/data/excelCatalog.ts').read_text(encoding='utf-8')
        return json.loads(text.split('export const excelCatalog = ', 1)[1].strip().removesuffix(';'))

    def test_correction_and_row_sort_retain_identity_membership_and_are_idempotent(self):
        self.run_import()
        before = self.catalog()
        book = openpyxl.load_workbook(self.source)
        words = book['Palabras']
        for row in (2, 3):
            words.cell(row, 4, 'buscar información')
        values = [cell.value for cell in words[2]]
        for column, cell in enumerate(words[3], 1):
            words.cell(2, column, cell.value)
        for column, value in enumerate(values, 1):
            words.cell(3, column, value)
        book.save(self.source)
        self.run_import()
        after = self.catalog()
        self.assertEqual(after['words'][0]['id'], before['words'][0]['id'])
        self.assertEqual(after['words'][0]['runtimeId'], before['words'][0]['runtimeId'])
        self.assertEqual(after['units'], before['units'])
        self.assertEqual(after['words'][0]['spanish'], 'buscar información')
        digest = hashlib.sha256(self.source.read_bytes()).hexdigest()
        output = (self.root / 'src/game/data/excelCatalog.ts').read_bytes()
        self.run_import()
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), digest)
        self.assertEqual((self.root / 'src/game/data/excelCatalog.ts').read_bytes(), output)

    def test_conflicting_duplicate_correction_fails_without_overwriting_catalog(self):
        self.run_import()
        before = self.catalog()
        book = openpyxl.load_workbook(self.source)
        book['Palabras'].cell(2, 4, 'buscar información')
        book.save(self.source)
        result = self.run_import(succeeds=False)
        self.assertIn('Conflicting text', result.stderr)
        self.assertEqual(self.catalog(), before)

    def test_new_rows_receive_persisted_ids_without_moving_old_races(self):
        self.run_import()
        before = self.catalog()
        book = openpyxl.load_workbook(self.source)
        book['Palabras'].append([1, 'Acciones', 'stand up', 'ponerse de pie', 'sleep', 'eat'])
        book.save(self.source)
        self.run_import()
        after = self.catalog()
        self.assertEqual(after['units'][0]['races'][0], before['units'][0]['races'][0])
        self.assertEqual(len(after['units'][0]['races']), 2)
        book = openpyxl.load_workbook(self.source)
        self.assertEqual(book['Palabras'].cell(4, 7).value, after['words'][1]['id'])
        self.assertEqual(book['Palabras'].cell(4, 8).value, after['words'][1]['runtimeId'])


if __name__ == '__main__':
    unittest.main()
