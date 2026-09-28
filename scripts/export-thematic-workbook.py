"""Export a reviewable thematic copy without overwriting an open source workbook."""
from copy import copy
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
CONTENT = ROOT / 'content/es-en'
text = (ROOT / 'src/game/data/thematicCatalog.ts').read_text(encoding='utf-8')
catalog = json.loads(text.split('export const thematicCatalog = ', 1)[1].strip().removesuffix(';'))
topics = {topic['id']: topic for topic in catalog['topics']}
book = openpyxl.load_workbook(CONTENT / 'espanol-ingles.xlsx')
sheet = book['Palabras']
for column, title in ((9, 'Categoría temática'), (10, 'Subtema'), (11, 'ID del subtema')):
    sheet.cell(1, column, title)._style = copy(sheet.cell(1, 2)._style)
    sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 32
for row in range(2, sheet.max_row + 1):
    wid = sheet.cell(row, 7).value
    topic = topics[catalog['wordTopics'][wid]]
    # A and Niveles are the original group IDs; keep them as historical metadata.
    sheet.cell(row, 2, topic['category'])
    for column, value in ((9, topic['category']), (10, topic['title']), (11, topic['id'])):
        cell = sheet.cell(row, column, value)
        cell._style = copy(sheet.cell(row, 2)._style)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
sheet.freeze_panes = 'C2'
sheet.auto_filter.ref = f'A1:K{sheet.max_row}'

def table(name, headers, widths):
    if name in book:
        del book[name]
    result = book.create_sheet(name)
    result.append(headers)
    result.freeze_panes = 'A2'
    for cell, width in zip(result[1], widths):
        cell.fill = PatternFill('solid', fgColor='176C50')
        cell.font = Font(name='Arial', size=11, color='FFFFFF', bold=True)
        result.column_dimensions[cell.column_letter].width = width
    return result

topic_sheet = table('Temas', ['Categoría', 'Subtema', 'ID estable'], [28, 36, 42])
race_sheet = table('Carreras temáticas', ['ID de carrera', 'Título', 'Categoría', 'Subtema', 'Palabras en inglés'], [54, 54, 28, 38, 90])
english = {sheet.cell(row, 7).value: sheet.cell(row, 3).value for row in range(2, sheet.max_row + 1)}
for topic in topics.values():
    topic_sheet.append([topic['category'], topic['title'], topic['id']])
    for race in topic['races']:
        race_sheet.append([race['id'], race['title'], topic['category'], topic['title'], ', '.join(english[wid] for wid in race['wordIds'])])
for result in (topic_sheet, race_sheet):
    for row in result.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        result.row_dimensions[row[0].row].height = 36
    result.auto_filter.ref = result.dimensions
notes = table('Leer categorías', ['Nota'], [110])
for note in [
    'Copia temática para revisar las categorías de Dashword; no reemplaza automáticamente el Excel editorial original.',
    'Palabras: C/D/E/F conservan las traducciones y distractores. G/H conservan los identificadores de aprendizaje.',
    'B e I muestran la nueva categoría. J contiene el subtema; K su identificador estable.',
    'A y la hoja Niveles conservan los grupos anteriores como referencia histórica. La nueva ruta está en Carreras temáticas.',
    'La clasificación editable del proyecto está en topic-lexicon.tsv y topic-senses.json. Regenerar con build-topic-catalog.py.',
    'Los resultados anteriores y los repasos permanecen guardados. Las nuevas carreras tienen sus propios resultados.',
]:
    notes.append([note])
    notes.cell(notes.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
    notes.cell(notes.max_row, 1).font = Font(name='Arial', size=11)
    notes.row_dimensions[notes.max_row].height = 32
output = CONTENT / 'espanol-ingles-categorizado.xlsx'
book.save(output)
print(output)
